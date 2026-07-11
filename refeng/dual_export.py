"""Dual-engine (Python + R) reference export (JSON, CSV, and optional Excel)."""

from __future__ import annotations

import csv
import hashlib
import json
import shutil
import subprocess
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from refeng.bootstrap import bootstrap_f2
from refeng.io import coefficient_of_variation, cv_gate_exceeded, load_dissolution_csv
from refeng.metrics import (
    compute_profile_metrics,
    ema_truncation,
    f1,
    f2,
    fda_truncation,
)
from refeng.models import MODEL_SPECS, fit_model_per_unit

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill(fill_type="solid", fgColor="FFD9E1F2")
NUM_FMT = "0.000000"

MODEL_LABELS = {
    "zero_order": "Zero-order",
    "first_order": "First-order",
    "higuchi": "Higuchi",
    "hixson_crowell": "Hixson-Crowell",
    "korsmeyer_peppas_full": "Korsmeyer-Peppas (full)",
    "korsmeyer_peppas_60": "Korsmeyer-Peppas (≤60%)",
    "weibull": "Weibull",
}

MODEL_PARAM_KEYS = {
    "zero_order": ["k0"],
    "first_order": ["k1"],
    "higuchi": ["kH"],
    "hixson_crowell": ["kHC"],
    "korsmeyer_peppas_full": ["kKP", "n"],
    "korsmeyer_peppas_60": ["kKP", "n"],
    "weibull": ["alpha", "beta"],
}

GOF_KEYS = ["R2", "adjR2", "RMSE", "AIC"]
ALL_PARAM_COLUMNS = ["k0", "k1", "kH", "kHC", "kKP", "n", "alpha", "beta"]

DATASETS: list[dict[str, str]] = [
    {
        "label": "polli_slow_similar",
        "ref": "polli_slow_similar_ref.csv",
        "test": "polli_slow_similar_test.csv",
        "time_unit": "min",
    },
    {
        "label": "polli_fast_vs_slow_dissimilar",
        "ref": "polli_fast_vs_slow_dissimilar_ref.csv",
        "test": "polli_fast_vs_slow_dissimilar_test.csv",
        "time_unit": "min",
    },
    {
        "label": "mendyk",
        "ref": "mendyk_reference_N12.csv",
        "test": "mendyk_test_N12.csv",
        "time_unit": "min",
    },
    {
        "label": "shah_similar",
        "ref": "shah_ref.csv",
        "test": "shah_test.csv",
        "time_unit": "min",
    },
    {
        "label": "shah_borderline",
        "ref": "shah_ref.csv",
        "test": "shah_borderline_test.csv",
        "time_unit": "min",
    },
    {
        "label": "tsong",
        "ref": "tsong_ref.csv",
        "test": "tsong_test.csv",
        "time_unit": "min",
    },
    {
        "label": "chow",
        "ref": "chow_ref.csv",
        "test": "chow_test.csv",
        "time_unit": "hours",
    },
]

BOOT_ITER = 10_000
BOOT_SEED = 42


@dataclass
class EngineDataset:
    label: str
    ref_file: str
    test_file: str
    time_unit: str
    n_timepoints: int
    n_ref_units: int
    n_test_units: int
    similarity: dict[str, Any]
    exposure: dict[str, float]
    bootstrap: dict[str, Any] | None
    ref_models: dict[str, dict]
    test_models: dict[str, dict]
    best_fit_ref: str | None


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    return digest.hexdigest()


def _abs_diff(a: float | None, b: float | None) -> float | None:
    if a is None or b is None:
        return None
    if isinstance(a, float) and np.isnan(a):
        return None
    if isinstance(b, float) and np.isnan(b):
        return None
    return abs(float(a) - float(b))


def _primary_metrics(
    times: np.ndarray, ref_units: np.ndarray, test_units: np.ndarray
) -> tuple[float, float, str]:
    ref_means = np.mean(ref_units, axis=0)
    test_means = np.mean(test_units, axis=0)
    fda = fda_truncation(times, ref_means, test_means)
    ema = ema_truncation(times, ref_means, test_means)
    if np.array_equal(fda.kept_indices, ema.kept_indices):
        trunc = fda
        rule = "FDA"
    else:
        trunc = ema
        rule = "EMA"
    return f2(trunc.ref_means, trunc.test_means), f1(trunc.ref_means, trunc.test_means), rule


def _first_point_cv(ref_units: np.ndarray, test_units: np.ndarray) -> float | None:
    ref_cv = coefficient_of_variation(ref_units)
    test_cv = coefficient_of_variation(test_units)
    ref_v = None if len(ref_cv) == 0 or np.isnan(ref_cv[0]) else float(ref_cv[0])
    test_v = None if len(test_cv) == 0 or np.isnan(test_cv[0]) else float(test_cv[0])
    if ref_v is None:
        return test_v
    if test_v is None:
        return ref_v
    return max(ref_v, test_v)


def _serialize_model(summary) -> dict[str, Any]:
    return {
        "converged": f"{summary.n_units_converged}/{summary.n_units_total}",
        "n_converged": summary.n_units_converged,
        "n_total": summary.n_units_total,
        "params": dict(summary.param_mean),
        "gof": {
            "R2": summary.gof_mean.get("R2"),
            "adjR2": summary.gof_mean.get("adj_R2"),
            "RMSE": summary.gof_mean.get("RMSE"),
            "AIC": summary.gof_mean.get("AIC"),
        },
    }


def _best_model_key(summaries: dict) -> str | None:
    best_key: str | None = None
    best_aic: float | None = None
    for key, summary in summaries.items():
        if key == "korsmeyer_peppas_60":
            continue
        if summary.n_units_converged <= 0.5 * summary.n_units_total:
            continue
        aic_val = summary.gof_mean.get("AIC")
        if aic_val is None:
            continue
        if best_aic is None or aic_val < best_aic:
            best_aic = aic_val
            best_key = key
    return best_key


def compute_python_dataset(data_dir: Path, spec: dict[str, str]) -> EngineDataset:
    ref_times, ref_units = load_dissolution_csv(data_dir / spec["ref"])
    test_times, test_units = load_dissolution_csv(data_dir / spec["test"])
    if not np.array_equal(ref_times, test_times):
        raise ValueError(f"{spec['label']}: reference and test time vectors must match")

    ref_models = {
        key: fit_model_per_unit(ref_times, ref_units, key) for key in MODEL_SPECS
    }
    test_models = {
        key: fit_model_per_unit(test_times, test_units, key) for key in MODEL_SPECS
    }

    f2_primary, f1_primary, primary_rule = _primary_metrics(ref_times, ref_units, test_units)
    metrics_fda = compute_profile_metrics(ref_times, ref_units, test_units, "FDA")

    ref_cv, _ = cv_gate_exceeded(ref_units)
    test_cv, _ = cv_gate_exceeded(test_units)
    high_cv = ref_cv or test_cv

    boot_block: dict[str, Any] | None = None
    if high_cv and spec["label"] == "mendyk":
        boot = bootstrap_f2(
            ref_times,
            ref_units,
            test_units,
            n_iterations=BOOT_ITER,
            seed=BOOT_SEED,
            truncation="EMA",
        )
        boot_block = {
            "E_f2": boot.e_f2,
            "ci_lower": boot.ci_lower,
            "ci_upper": boot.ci_upper,
            "decision": "similar" if boot.similar else "not similar",
        }

    return EngineDataset(
        label=spec["label"],
        ref_file=spec["ref"],
        test_file=spec["test"],
        time_unit=spec["time_unit"],
        n_timepoints=len(ref_times),
        n_ref_units=ref_units.shape[0],
        n_test_units=test_units.shape[0],
        similarity={
            "f2_primary": f2_primary,
            "f1_primary": f1_primary,
            "primary_rule": primary_rule,
            "first_point_CV": _first_point_cv(ref_units, test_units),
            "high_variability": high_cv,
        },
        exposure={
            "DE_ref": metrics_fda["DE_ref"],
            "DE_test": metrics_fda["DE_test"],
            "MDT_ref": metrics_fda["MDT_ref"],
            "MDT_test": metrics_fda["MDT_test"],
        },
        bootstrap=boot_block,
        ref_models={k: _serialize_model(v) for k, v in ref_models.items()},
        test_models={k: _serialize_model(v) for k, v in test_models.items()},
        best_fit_ref=_best_model_key(ref_models),
    )


def _parse_r_dataset(raw: dict) -> EngineDataset:
    def _models(block: dict) -> dict[str, dict]:
        out: dict[str, dict] = {}
        for key, value in block.items():
            params = value.get("params")
            if isinstance(params, list) and len(params) == 0:
                params = {}
            out[key] = {
                **value,
                "params": params or {},
            }
        return out

    return EngineDataset(
        label=raw["label"],
        ref_file=raw["ref_file"],
        test_file=raw["test_file"],
        time_unit=raw["time_unit"],
        n_timepoints=int(raw["n_timepoints"]),
        n_ref_units=int(raw["n_ref_units"]),
        n_test_units=int(raw["n_test_units"]),
        similarity=raw["similarity"],
        exposure=raw["exposure"],
        bootstrap=raw.get("bootstrap"),
        ref_models=_models(raw["ref_models"]),
        test_models=_models(raw["test_models"]),
        best_fit_ref=raw.get("best_fit_ref"),
    )


def run_r_engine(data_dir: Path, script_dir: Path) -> tuple[list[EngineDataset], dict[str, str]]:
    rscript = shutil.which("Rscript")
    if rscript is None:
        raise RuntimeError("Rscript not found; install R with minpack.lm")

    script = script_dir / "export_r_reference_json.R"
    reference_r = script_dir.parent / "reference.R"
    proc = subprocess.run(
        [rscript, str(script), str(data_dir), str(reference_r)],
        capture_output=True,
        text=True,
        check=False,
    )
    if proc.returncode != 0:
        raise RuntimeError(
            f"R export failed (exit {proc.returncode}):\n{proc.stderr or proc.stdout}"
        )
    payload = json.loads(proc.stdout)
    versions = payload.get("versions", {})
    return [_parse_r_dataset(item) for item in payload["datasets"]], versions


def _python_versions() -> dict[str, str]:
    import numpy
    import scipy

    return {
        "python": f"{sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro}",
        "numpy": numpy.__version__,
        "scipy": scipy.__version__,
    }


def _style_sheet(ws) -> None:
    if ws.max_row < 1:
        return
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(
            min_col=col_idx, max_col=col_idx, min_row=1, max_row=ws.max_row
        ):
            val = row[0].value
            max_len = max(max_len, len(str(val)) if val is not None else 0)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 60)


def _write_row(ws, row_idx: int, values: list[Any], numeric_cols: set[int] | None = None) -> None:
    numeric_cols = numeric_cols or set()
    for col_idx, val in enumerate(values, start=1):
        if isinstance(val, dict):
            val = None
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        if col_idx in numeric_cols and isinstance(val, (int, float)):
            cell.number_format = NUM_FMT


def _triplet_headers(base: str) -> list[str]:
    return [f"{base}[R]", f"{base}[Py]", f"{base}_AbsDiff"]


def _write_readme(
    ws,
    data_dir: Path,
    py_versions: dict[str, str],
    r_versions: dict[str, str],
) -> None:
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    rows: list[list[Any]] = [
        ["Field", "Value"],
        ["Export date (UTC)", now],
        ["Python engine", "scipy.optimize.curve_fit (TRF/LM)"],
        ["R engine", "minpack.lm::nlsLM"],
        ["Python version", py_versions["python"]],
        ["numpy version", py_versions["numpy"]],
        ["scipy version", py_versions["scipy"]],
        ["R version", r_versions.get("R", "")],
        ["minpack.lm version", r_versions.get("minpack_lm", "")],
        ["bootf2 version", r_versions.get("bootf2", "")],
        ["Bootstrap", f"unit-level resampling, seed={BOOT_SEED}, n={BOOT_ITER}, EMA-primary"],
        ["", ""],
        [
            "Dataset",
            "Ref CSV",
            "Test CSV",
            "sha256(ref)",
            "sha256(test)",
            "Ref units",
            "Test units",
            "Timepoints",
            "Time unit",
        ],
    ]
    for spec in DATASETS:
        ref_path = data_dir / spec["ref"]
        test_path = data_dir / spec["test"]
        _, ref_units = load_dissolution_csv(ref_path)
        _, test_units = load_dissolution_csv(test_path)
        times, _ = load_dissolution_csv(ref_path)
        rows.append(
            [
                spec["label"],
                spec["ref"],
                spec["test"],
                _sha256(ref_path),
                _sha256(test_path),
                ref_units.shape[0],
                test_units.shape[0],
                len(times),
                spec["time_unit"],
            ]
        )
    for i, row in enumerate(rows, start=1):
        _write_row(ws, i, row)
    _style_sheet(ws)


def _write_similarity(ws, py_by_label: dict, r_by_label: dict) -> None:
    headers = ["Dataset", "primary_rule"]
    for metric in ("f2_primary", "f1_primary", "first_point_CV"):
        headers.extend(_triplet_headers(metric))
    boot_metrics = ("E_f2", "ci_lower", "ci_upper", "decision")
    for metric in boot_metrics:
        headers.extend([f"{metric}[R]", f"{metric}[Py]"])
        if metric != "decision":
            headers.append(f"{metric}_AbsDiff")
    headers.append("high_variability")

    rows: list[list[Any]] = []
    numeric_cols: set[int] = set()
    for spec in DATASETS:
        label = spec["label"]
        py = py_by_label[label]
        r = r_by_label[label]
        row: list[Any] = [label, py.similarity.get("primary_rule")]
        col = 3
        for metric in ("f2_primary", "f1_primary", "first_point_CV"):
            rv = r.similarity.get(metric)
            pv = py.similarity.get(metric)
            row.extend([rv, pv, _abs_diff(rv, pv)])
            numeric_cols.update({col, col + 1, col + 2})
            col += 3
        high_cv = bool(py.similarity.get("high_variability"))
        if high_cv:
            for metric in boot_metrics:
                rb = (r.bootstrap or {}).get(metric)
                pb = (py.bootstrap or {}).get(metric)
                row.extend([rb, pb])
                col += 2
                if metric != "decision":
                    row.append(_abs_diff(rb, pb) if metric != "decision" else None)
                    numeric_cols.add(col)
                    col += 1
        else:
            for metric in boot_metrics:
                row.extend([None, None])
                col += 2
                if metric != "decision":
                    row.append(None)
                    col += 1
        row.append("yes" if high_cv else "no")
        rows.append(row)

    _write_row(ws, 1, headers)
    for i, row in enumerate(rows, start=2):
        _write_row(ws, i, row, numeric_cols)
    _style_sheet(ws)


def _write_exposure(ws, py_by_label: dict, r_by_label: dict) -> None:
    headers = ["Dataset", "time_unit"]
    for metric in ("DE_ref", "DE_test", "MDT_ref", "MDT_test"):
        headers.extend(_triplet_headers(metric))
    rows: list[list[Any]] = []
    numeric_cols: set[int] = set()
    for spec in DATASETS:
        label = spec["label"]
        py = py_by_label[label]
        r = r_by_label[label]
        row: list[Any] = [label, spec["time_unit"]]
        col = 3
        for metric in ("DE_ref", "DE_test", "MDT_ref", "MDT_test"):
            rv = r.exposure.get(metric)
            pv = py.exposure.get(metric)
            row.extend([rv, pv, _abs_diff(rv, pv)])
            numeric_cols.update({col, col + 1, col + 2})
            col += 3
        rows.append(row)
    _write_row(ws, 1, headers)
    for i, row in enumerate(rows, start=2):
        _write_row(ws, i, row, numeric_cols)
    _style_sheet(ws)


def _model_value(block: dict | None, field: str) -> Any:
    if not block:
        return None
    if field == "converged":
        return block.get("converged")
    if field.startswith("gof."):
        return (block.get("gof") or {}).get(field.split(".", 1)[1])
    if field.startswith("params."):
        return (block.get("params") or {}).get(field.split(".", 1)[1])
    return None


def _write_model_fits(ws, py_by_label: dict, r_by_label: dict) -> None:
    headers = ["Dataset", "Profile", "Model", "converged[R]", "converged[Py]"]
    for gof in GOF_KEYS:
        headers.extend(_triplet_headers(gof))
    for param in ALL_PARAM_COLUMNS:
        headers.extend(_triplet_headers(param))

    rows: list[list[Any]] = []
    numeric_cols: set[int] = set()

    for spec in DATASETS:
        label = spec["label"]
        py = py_by_label[label]
        r = r_by_label[label]
        for profile, py_models, r_models in (
            ("ref", py.ref_models, r.ref_models),
            ("test", py.test_models, r.test_models),
        ):
            for model_key, model_label in MODEL_LABELS.items():
                py_block = py_models.get(model_key)
                r_block = r_models.get(model_key)
                row: list[Any] = [label, profile, model_label]
                col = 4
                rv = _model_value(r_block, "converged")
                pv = _model_value(py_block, "converged")
                row.extend([rv, pv])
                col += 2
                for gof in GOF_KEYS:
                    rv = _model_value(r_block, f"gof.{gof}")
                    pv = _model_value(py_block, f"gof.{gof}")
                    row.extend([rv, pv, _abs_diff(rv, pv)])
                    numeric_cols.update({col, col + 1, col + 2})
                    col += 3
                for param in ALL_PARAM_COLUMNS:
                    rv = _model_value(r_block, f"params.{param}")
                    pv = _model_value(py_block, f"params.{param}")
                    row.extend([rv, pv, _abs_diff(rv, pv)])
                    if isinstance(rv, (int, float)) or isinstance(pv, (int, float)):
                        numeric_cols.update({col, col + 1, col + 2})
                    col += 3
                rows.append(row)

    _write_row(ws, 1, headers)
    for i, row in enumerate(rows, start=2):
        _write_row(ws, i, row, numeric_cols)
    _style_sheet(ws)


def _write_best_fit(ws, py_by_label: dict, r_by_label: dict) -> None:
    headers = [
        "Dataset",
        "best_fit_model[R]",
        "best_fit_model[Py]",
        "models_agree",
    ]
    rows: list[list[Any]] = []
    for spec in DATASETS:
        label = spec["label"]
        r_key = r_by_label[label].best_fit_ref
        py_key = py_by_label[label].best_fit_ref
        r_name = MODEL_LABELS.get(r_key, r_key) if r_key else None
        py_name = MODEL_LABELS.get(py_key, py_key) if py_key else None
        agree = "yes" if r_key is not None and r_key == py_key else "no"
        rows.append([label, r_name, py_name, agree])
    _write_row(ws, 1, headers)
    for i, row in enumerate(rows, start=2):
        _write_row(ws, i, row)
    _style_sheet(ws)


def write_reference_results_workbook(
    output_path: Path,
    data_dir: Path,
    script_dir: Path,
) -> Path:
    py_versions = _python_versions()
    py_datasets = [compute_python_dataset(data_dir, spec) for spec in DATASETS]
    r_payload_datasets, r_versions = run_r_engine(data_dir, script_dir)

    py_by_label = {d.label: d for d in py_datasets}
    r_by_label = {d.label: d for d in r_payload_datasets}

    output_path = output_path.resolve()
    wb = Workbook()

    ws_readme = wb.active
    ws_readme.title = "README"
    _write_readme(ws_readme, data_dir, py_versions, r_versions)

    ws_sim = wb.create_sheet("Similarity")
    _write_similarity(ws_sim, py_by_label, r_by_label)

    ws_exp = wb.create_sheet("Exposure")
    _write_exposure(ws_exp, py_by_label, r_by_label)

    ws_models = wb.create_sheet("ModelFits")
    _write_model_fits(ws_models, py_by_label, r_by_label)

    ws_best = wb.create_sheet("BestFit")
    _write_best_fit(ws_best, py_by_label, r_by_label)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def _dataset_to_dict(d: EngineDataset) -> dict[str, Any]:
    return {
        "label": d.label,
        "ref_file": d.ref_file,
        "test_file": d.test_file,
        "time_unit": d.time_unit,
        "n_timepoints": d.n_timepoints,
        "n_ref_units": d.n_ref_units,
        "n_test_units": d.n_test_units,
        "similarity": d.similarity,
        "exposure": d.exposure,
        "bootstrap": d.bootstrap,
        "ref_models": d.ref_models,
        "test_models": d.test_models,
        "best_fit_ref": d.best_fit_ref,
    }


def _build_export_payload(
    data_dir: Path,
    script_dir: Path,
) -> tuple[dict[str, EngineDataset], dict[str, EngineDataset], dict[str, str], dict[str, str]]:
    py_versions = _python_versions()
    py_datasets = [compute_python_dataset(data_dir, spec) for spec in DATASETS]
    r_payload_datasets, r_versions = run_r_engine(data_dir, script_dir)
    py_by_label = {d.label: d for d in py_datasets}
    r_by_label = {d.label: d for d in r_payload_datasets}
    return py_by_label, r_by_label, py_versions, r_versions


def _similarity_csv_rows(py_by_label: dict, r_by_label: dict) -> tuple[list[str], list[list[Any]]]:
    headers = ["dataset", "primary_rule"]
    for metric in ("f2_primary", "f1_primary", "first_point_CV"):
        headers.extend(_triplet_headers(metric))
    boot_metrics = ("E_f2", "ci_lower", "ci_upper", "decision")
    for metric in boot_metrics:
        headers.extend([f"{metric}[R]", f"{metric}[Py]"])
        if metric != "decision":
            headers.append(f"{metric}_AbsDiff")
    headers.append("high_variability")

    rows: list[list[Any]] = []
    for spec in DATASETS:
        label = spec["label"]
        py = py_by_label[label]
        r = r_by_label[label]
        row: list[Any] = [label, py.similarity.get("primary_rule")]
        for metric in ("f2_primary", "f1_primary", "first_point_CV"):
            rv = r.similarity.get(metric)
            pv = py.similarity.get(metric)
            row.extend([rv, pv, _abs_diff(rv, pv)])
        high_cv = bool(py.similarity.get("high_variability"))
        if high_cv:
            for metric in boot_metrics:
                rb = (r.bootstrap or {}).get(metric)
                pb = (py.bootstrap or {}).get(metric)
                row.extend([rb, pb])
                if metric != "decision":
                    row.append(_abs_diff(rb, pb) if metric != "decision" else None)
        else:
            for metric in boot_metrics:
                row.extend([None, None])
                if metric != "decision":
                    row.append(None)
        row.append("yes" if high_cv else "no")
        rows.append(row)
    return headers, rows


def _exposure_csv_rows(py_by_label: dict, r_by_label: dict) -> tuple[list[str], list[list[Any]]]:
    headers = ["dataset", "time_unit"]
    for metric in ("DE_ref", "DE_test", "MDT_ref", "MDT_test"):
        headers.extend(_triplet_headers(metric))
    rows: list[list[Any]] = []
    for spec in DATASETS:
        label = spec["label"]
        py = py_by_label[label]
        r = r_by_label[label]
        row: list[Any] = [label, spec["time_unit"]]
        for metric in ("DE_ref", "DE_test", "MDT_ref", "MDT_test"):
            rv = r.exposure.get(metric)
            pv = py.exposure.get(metric)
            row.extend([rv, pv, _abs_diff(rv, pv)])
        rows.append(row)
    return headers, rows


def _best_fit_csv_rows(py_by_label: dict, r_by_label: dict) -> tuple[list[str], list[list[Any]]]:
    headers = ["dataset", "best_fit_model[R]", "best_fit_model[Py]", "models_agree"]
    rows: list[list[Any]] = []
    for spec in DATASETS:
        label = spec["label"]
        r_key = r_by_label[label].best_fit_ref
        py_key = py_by_label[label].best_fit_ref
        r_name = MODEL_LABELS.get(r_key, r_key) if r_key else None
        py_name = MODEL_LABELS.get(py_key, py_key) if py_key else None
        agree = "yes" if r_key is not None and r_key == py_key else "no"
        rows.append([label, r_name, py_name, agree])
    return headers, rows


def _model_fits_csv_rows(py_by_label: dict, r_by_label: dict) -> tuple[list[str], list[list[Any]]]:
    headers = ["dataset", "profile", "model", "converged[R]", "converged[Py]"]
    for gof in GOF_KEYS:
        headers.extend(_triplet_headers(gof))
    for param in ALL_PARAM_COLUMNS:
        headers.extend(_triplet_headers(param))

    rows: list[list[Any]] = []
    for spec in DATASETS:
        label = spec["label"]
        py = py_by_label[label]
        r = r_by_label[label]
        for profile, py_models, r_models in (
            ("ref", py.ref_models, r.ref_models),
            ("test", py.test_models, r.test_models),
        ):
            for model_key, model_label in MODEL_LABELS.items():
                py_block = py_models.get(model_key)
                r_block = r_models.get(model_key)
                row: list[Any] = [label, profile, model_label]
                row.extend([_model_value(r_block, "converged"), _model_value(py_block, "converged")])
                for gof in GOF_KEYS:
                    rv = _model_value(r_block, f"gof.{gof}")
                    pv = _model_value(py_block, f"gof.{gof}")
                    row.extend([rv, pv, _abs_diff(rv, pv)])
                for param in ALL_PARAM_COLUMNS:
                    rv = _model_value(r_block, f"params.{param}")
                    pv = _model_value(py_block, f"params.{param}")
                    row.extend([rv, pv, _abs_diff(rv, pv)])
                rows.append(row)
    return headers, rows


def _write_csv(path: Path, headers: list[str], rows: list[list[Any]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers)
        writer.writerows(rows)
    return path


def write_reference_results_exports(
    output_dir: Path,
    data_dir: Path,
    script_dir: Path,
) -> dict[str, Path]:
    """Write dual-engine comparison as JSON and flat CSV tables."""
    py_by_label, r_by_label, py_versions, r_versions = _build_export_payload(data_dir, script_dir)
    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    dataset_manifest = []
    for spec in DATASETS:
        ref_path = data_dir / spec["ref"]
        test_path = data_dir / spec["test"]
        _, ref_units = load_dissolution_csv(ref_path)
        _, test_units = load_dissolution_csv(test_path)
        times, _ = load_dissolution_csv(ref_path)
        dataset_manifest.append(
            {
                "label": spec["label"],
                "ref_file": spec["ref"],
                "test_file": spec["test"],
                "sha256_ref": _sha256(ref_path),
                "sha256_test": _sha256(test_path),
                "n_ref_units": ref_units.shape[0],
                "n_test_units": test_units.shape[0],
                "n_timepoints": len(times),
                "time_unit": spec["time_unit"],
            }
        )

    payload = {
        "export_date_utc": now,
        "engines": {
            "python": {
                "implementation": "scipy.optimize.curve_fit (TRF/LM)",
                "versions": py_versions,
            },
            "r": {
                "implementation": "minpack.lm::nlsLM",
                "versions": r_versions,
            },
        },
        "bootstrap": {
            "method": "unit-level resampling",
            "seed": BOOT_SEED,
            "iterations": BOOT_ITER,
            "truncation": "EMA-primary when high variability (mendyk only)",
        },
        "datasets": dataset_manifest,
        "python": {label: _dataset_to_dict(py_by_label[label]) for label in py_by_label},
        "r": {label: _dataset_to_dict(r_by_label[label]) for label in r_by_label},
    }

    json_path = output_dir / "reference_results.json"
    json_path.write_text(json.dumps(payload, indent=2) + "\n")

    sim_h, sim_r = _similarity_csv_rows(py_by_label, r_by_label)
    exp_h, exp_r = _exposure_csv_rows(py_by_label, r_by_label)
    mod_h, mod_r = _model_fits_csv_rows(py_by_label, r_by_label)
    best_h, best_r = _best_fit_csv_rows(py_by_label, r_by_label)

    return {
        "json": json_path,
        "similarity": _write_csv(output_dir / "similarity.csv", sim_h, sim_r),
        "exposure": _write_csv(output_dir / "exposure.csv", exp_h, exp_r),
        "model_fits": _write_csv(output_dir / "model_fits.csv", mod_h, mod_r),
        "best_fit": _write_csv(output_dir / "best_fit.csv", best_h, best_r),
    }
