"""Pytest suite for the independent dissolution reference engine."""

from __future__ import annotations

from pathlib import Path

import json
import numpy as np
import pytest

from refeng.polli_sets import POLLI_SLOW_SIMILAR

ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture(scope="module")
def golden_polli_slow_similar() -> dict:
    path = ROOT / "golden" / "golden_polli_slow_similar.json"
    with path.open() as fh:
        return json.load(fh)


@pytest.fixture(scope="module")
def polli_slow_similar_data():
    from refeng.io import load_dissolution_csv

    times, ref_units = load_dissolution_csv(
        ROOT / "data" / POLLI_SLOW_SIMILAR.ref_file
    )
    _, test_units = load_dissolution_csv(
        ROOT / "data" / POLLI_SLOW_SIMILAR.test_file
    )
    return times, ref_units, test_units


def test_polli_slow_similar_csv_loading(polli_slow_similar_data):
    times, ref_units, test_units = polli_slow_similar_data
    assert len(times) == 7
    assert ref_units.shape == (6, 7)
    assert test_units.shape == (1, 7)


def test_polli_slow_similar_golden_model_fits(
    polli_slow_similar_data, golden_polli_slow_similar
):
    from refeng.models import fit_model_per_unit

    times, ref_units, _ = polli_slow_similar_data
    expected = golden_polli_slow_similar["polli_slow_similar"]["reference_model_fits"]
    tol = golden_polli_slow_similar["_meta"]["tolerance_deterministic"]

    weibull = fit_model_per_unit(times, ref_units, "weibull")
    w = expected["Weibull"]
    assert abs(weibull.param_mean["alpha"] - w["alpha"]) < tol
    assert abs(weibull.param_mean["beta"] - w["beta"]) < tol
    assert abs(weibull.gof_mean["R2"] - w["R2"]) < tol
    assert abs(weibull.gof_mean["AIC"] - w["AIC"]) < tol
    assert abs(weibull.gof_mean["RMSE"] - w["RMSE"]) < tol

    fo = fit_model_per_unit(times, ref_units, "first_order")
    f = expected["First-order"]
    assert abs(fo.param_mean["k1"] - f["k"]) < tol
    assert abs(fo.gof_mean["R2"] - f["R2"]) < tol
    assert abs(fo.gof_mean["AIC"] - f["AIC"]) < tol


def test_polli_slow_similar_golden_similarity_f2_f1(
    polli_slow_similar_data, golden_polli_slow_similar
):
    from refeng.metrics import compute_profile_metrics

    times, ref_units, test_units = polli_slow_similar_data
    metrics = compute_profile_metrics(times, ref_units, test_units, truncation="FDA")
    sim = golden_polli_slow_similar["polli_slow_similar"]["similarity_ref_vs_test_unit6"]
    tol = golden_polli_slow_similar["_meta"]["tolerance_deterministic"]

    assert abs(metrics["f2"] - sim["f2"]) < tol
    assert abs(metrics["f1"] - sim["f1"]) < tol
    assert metrics["truncated_times"] == sim["truncation_points_min"]
    assert abs(metrics["DE_ref"] - sim["DE_ref"]) < tol
    assert abs(metrics["DE_test"] - sim["DE_test"]) < tol
    assert abs(metrics["MDT_ref"] - sim["MDT_ref"]) < tol
    assert abs(metrics["MDT_test"] - sim["MDT_test"]) < tol


def test_single_unit_test_matrix_bootstrap(polli_slow_similar_data):
    from refeng.bootstrap import bootstrap_f2, resample_unit_indices
    from refeng.io import cv_gate_exceeded

    times, ref_units, test_units = polli_slow_similar_data
    flag, _ = cv_gate_exceeded(test_units)
    assert flag is False
    boot = bootstrap_f2(times, ref_units, test_units, n_iterations=100, seed=1)
    assert np.isfinite(boot.e_f2)
    assert np.isfinite(boot.ci_lower)
    rng = np.random.default_rng(1)
    for _ in range(20):
        idx = resample_unit_indices(test_units.shape[0], rng)
        assert np.allclose(test_units[idx], test_units[0])


def test_gof_formulas():
    from refeng.gof import aic, adjusted_r_squared, msc, r_squared, rmse

    observed = np.array([10.0, 25.0, 45.0, 70.0, 88.0])
    predicted = np.array([12.0, 24.0, 46.0, 68.0, 90.0])
    n_params = 2

    r2 = r_squared(observed, predicted)
    assert 0.0 < r2 < 1.0
    adj = adjusted_r_squared(observed, predicted, n_params)
    assert adj is not None and adj < r2
    assert rmse(observed, predicted, n_params) is not None
    assert rmse(observed, predicted, n_params) > 0
    assert aic(observed, predicted, n_params) > 0
    assert isinstance(msc(observed, predicted, n_params), float)


def test_fda_and_ema_truncation():
    from refeng.metrics import ema_truncation, fda_truncation, f1, f2

    times = np.array([10.0, 20.0, 30.0, 45.0, 60.0])
    ref = np.array([20.0, 40.0, 60.0, 86.0, 92.0])
    test = np.array([22.0, 42.0, 62.0, 84.0, 90.0])

    fda = fda_truncation(times, ref, test)
    assert fda.kept_indices[-1] == 4

    ema = ema_truncation(times, ref, test)
    assert ema.kept_indices[-1] == 3

    f2_fda = f2(fda.ref_means, fda.test_means)
    f1_fda = f1(fda.ref_means, fda.test_means)
    assert f2_fda > 50
    assert f1_fda >= 0


def test_bootstrap_is_unit_level(polli_slow_similar_data):
    from refeng.bootstrap import resample_unit_indices

    _, ref_units, test_units = polli_slow_similar_data
    rng = np.random.default_rng(123)

    for _ in range(20):
        ref_idx = resample_unit_indices(ref_units.shape[0], rng)
        test_idx = resample_unit_indices(test_units.shape[0], rng)
        ref_sample = ref_units[ref_idx]
        test_sample = test_units[test_idx]
        for i, ridx in enumerate(ref_idx):
            assert np.allclose(ref_sample[i], ref_units[ridx])
        for i, tidx in enumerate(test_idx):
            assert np.allclose(test_sample[i], test_units[tidx])


def test_cv_gate_mendyk():
    from refeng.io import cv_gate_exceeded, load_dissolution_csv

    _, ref = load_dissolution_csv(ROOT / "data" / "mendyk_reference_N12.csv")
    flag, msg = cv_gate_exceeded(ref)
    assert flag is True
    assert msg


def test_cli_golden_polli_slow_similar():
    import subprocess
    import sys

    cmd = [
        sys.executable,
        str(ROOT / "run_benchmark.py"),
        "--data",
        str(ROOT / "data"),
        "--pairs",
        "polli_slow_similar_ref.csv:polli_slow_similar_test.csv",
        "--golden",
        str(ROOT / "golden" / "golden_polli_slow_similar.json"),
        "--bootstrap-iterations",
        "200",
    ]
    env = {**subprocess.os.environ, "PYTHONPATH": str(ROOT)}
    proc = subprocess.run(cmd, capture_output=True, text=True, env=env)
    out = proc.stdout + proc.stderr
    assert "GoldenCheck" in out
    assert proc.returncode == 0, out
    assert "polli_slow_similar" in out
    assert "FAIL" not in out.split("GoldenCheck")[-1]


def test_cli_excel_export(tmp_path):
    import subprocess
    import sys

    from openpyxl import load_workbook

    out_file = tmp_path / "benchmark_out.xlsx"
    cmd = [
        sys.executable,
        str(ROOT / "run_benchmark.py"),
        "--data",
        str(ROOT / "data"),
        "--pairs",
        "polli_slow_similar_ref.csv:polli_slow_similar_test.csv,mendyk_reference_N12.csv:mendyk_test_N12.csv",
        "--golden",
        str(ROOT / "golden" / "golden_polli_slow_similar.json"),
        "--excel",
        str(out_file),
        "--bootstrap-iterations",
        "200",
    ]
    env = {**subprocess.os.environ, "PYTHONPATH": str(ROOT)}
    proc = subprocess.run(cmd, capture_output=True, text=True, env=env)
    assert proc.returncode in (0, 1), proc.stdout + proc.stderr
    assert out_file.exists()

    wb = load_workbook(out_file)
    ws = wb["Summary"]
    assert ws.max_row == 1 + 2
    assert ws["A2"].value == "polli_slow_similar_ref_vs_polli_slow_similar_test"
    golden_ws = wb["GoldenCheck"]
    assert golden_ws.max_row >= 2
    wb.close()


def test_polli_fast_vs_slow_dissimilar_not_golden():
    """Fast-vs-slow must not be checked against polli_slow_similar golden."""
    from refeng.golden_check import compare_polli_slow_similar_pair
    from refeng.excel_export import PairExportData

    pair = PairExportData(
        pair_name="polli_fast_vs_slow_dissimilar",
        ref_file="polli_fast_vs_slow_dissimilar_ref.csv",
        test_file="polli_fast_vs_slow_dissimilar_test.csv",
        metrics_fda={},
        metrics_ema={},
    )
    with open(ROOT / "golden" / "golden_polli_slow_similar.json") as fh:
        golden = json.load(fh)
    assert compare_polli_slow_similar_pair(pair, golden) == []
